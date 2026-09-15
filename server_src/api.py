"""API routes — REST endpoints for keypair, servers, assessment, config, and reports."""
import uuid
import json
import logging
from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from fastapi.responses import JSONResponse, FileResponse, HTMLResponse
from pydantic import BaseModel
from typing import Optional, List
import openpyxl
import io

from ..config import state, BENCHMARK_DIR, REPORT_DIR, KEYPAIR_DIR, AVAILABLE_MODELS
from ..services.keypair_service import generate_keypair, get_public_key_instructions, save_keypair_to_disk, load_keypair_from_disk

from ..services.ssh_service import SSHService
from ..services.rag_service import RAGService
from ..services.llm_service import LLMService
from ..services.assessment_engine import AssessmentEngine
from ..services.report_service import ReportService

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api")

# Initialize RAG service
rag_service = RAGService(BENCHMARK_DIR)
report_service = ReportService(REPORT_DIR)


# ============================================================
# Keypair endpoints
# ============================================================

class KeypairResponse(BaseModel):
    public_key: str
    fingerprint: str
    instructions: str


@router.post("/keypair/generate")
async def generate_keypair_api():
    """Generate a new SSH keypair. Saved to disk for persistence."""
    try:
        result = generate_keypair()
        save_keypair_to_disk(result, KEYPAIR_DIR)
        # Store in memory: (private_key_obj, public_key_str)
        state.keypair = (result["private_key"], result["public_key"])
        instructions = get_public_key_instructions(result["public_key"])

        return {
            "success": True,
            "public_key": result["public_key"],
            "fingerprint": result["fingerprint"],
            "instructions": instructions,
        }
    except Exception as e:
        logger.error(f"Keypair generation failed: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/keypair/status")
async def keypair_status():
    """Check if keypair exists."""
    return {
        "exists": state.keypair is not None,
        "public_key": state.keypair[1] if state.keypair else None,
    }


# ============================================================
# Server management endpoints
# ============================================================

class ServerConfig(BaseModel):
    ip: str
    ssh_port: int = 22
    ssh_username: str = "root"


@router.post("/servers/add")
async def add_server(server: ServerConfig):
    """Add a server for assessment."""
    server_id = str(uuid.uuid4())[:8]
    server_dict = {
        "id": server_id,
        **server.dict(),
    }
    state.servers.append(server_dict)
    return {"success": True, "server_id": server_id, "server": server_dict}


@router.post("/servers/add-batch")
async def add_servers_batch(servers: List[ServerConfig]):
    """Add multiple servers at once."""
    added = []
    for server in servers:
        server_id = str(uuid.uuid4())[:8]
        server_dict = {"id": server_id, **server.dict()}
        state.servers.append(server_dict)
        added.append(server_dict)
    return {"success": True, "count": len(added), "servers": added}


@router.get("/servers/list")
async def list_servers():
    """List all configured servers."""
    return {"servers": state.servers, "count": len(state.servers)}


@router.delete("/servers/{server_id}")
async def remove_server(server_id: str):
    """Remove a server."""
    before = len(state.servers)
    state.servers = [s for s in state.servers if s["id"] != server_id]
    if len(state.servers) == before:
        raise HTTPException(status_code=404, detail="Server not found")
    return {"success": True}


class TestConnectionRequest(BaseModel):
    ip: str
    ssh_port: int = 22
    ssh_username: str = "root"


@router.post("/servers/test-connection")
async def test_ssh_connection(req: TestConnectionRequest):
    """Test SSH connection to a server before running assessment."""
    if not state.keypair:
        return {"success": False, "error": "Keypair not generated. Please generate a keypair first."}

    try:
        private_key = state.keypair[0]
        ssh_service = SSHService(private_key, timeout=10)
        client = ssh_service.connect(req.ip, req.ssh_port, req.ssh_username)

        # Run a simple test command
        result = ssh_service.execute_command(client, "echo SSH_CONNECTION_OK && uname -a")
        client.close()

        if result["exit_status"] == 0:
            return {
                "success": True,
                "message": f"SSH connection to {req.ip}:{req.ssh_port} successful!",
                "server_info": result["output"],
            }
        else:
            return {
                "success": False,
                "error": f"SSH connected but test command failed: {result.get('error', 'unknown')}",
            }
    except Exception as e:
        return {"success": False, "error": f"SSH connection failed: {str(e)}"}


# ============================================================
# Benchmark visibility endpoints
# ============================================================

@router.get("/benchmarks/list")
async def list_benchmarks():
    """List all available benchmarks with their control counts."""
    benchmarks = []
    benchmark_files = {
        "CIS": "cis_linux.json",
        "STIG": "stig_linux.json",
        "NIST": "nist_linux.json",
        "PCI": "pci_linux.json",
    }

    for benchmark_type, filename in benchmark_files.items():
        filepath = BENCHMARK_DIR / filename
        if filepath.exists():
            with open(filepath, "r", encoding="utf-8") as f:
                data = json.load(f)
            controls = data.get("controls", [])
            categories = list(set(c.get("category", "Unknown") for c in controls))
            severities = {}
            for c in controls:
                sev = c.get("severity", "Unknown")
                severities[sev] = severities.get(sev, 0) + 1
            benchmarks.append({
                "benchmark": benchmark_type,
                "platform": data.get("platform", "Linux"),
                "version": data.get("version", ""),
                "description": data.get("description", ""),
                "control_count": len(controls),
                "categories": sorted(categories),
                "severity_summary": severities,
            })

    return {"benchmarks": benchmarks}


@router.get("/benchmarks/{benchmark_type}")
async def get_benchmark_controls(benchmark_type: str):
    """Get all controls in a specific benchmark."""
    benchmark_type = benchmark_type.upper()

    filename = f"{benchmark_type.lower()}_linux.json"
    filepath = BENCHMARK_DIR / filename

    if not filepath.exists():
        raise HTTPException(status_code=404, detail=f"Benchmark file not found: {filename}")

    with open(filepath, "r", encoding="utf-8") as f:
        data = json.load(f)

    return {
        "benchmark": benchmark_type,
        "platform": data.get("platform", "Linux"),
        "version": data.get("version", ""),
        "description": data.get("description", ""),
        "control_count": len(data.get("controls", [])),
        "controls": data.get("controls", []),
    }


# ============================================================
# LLM configuration endpoints
# ============================================================

class LLMConfigUpdate(BaseModel):
    api_key: str
    base_url: str = "https://api-ap-southeast-1.modelarts-maas.com/v2"
    model: str = "glm-5.2"
    temperature: float = 0.1
    rag_enabled: bool = True



@router.post("/config/llm")
async def update_llm_config(config: LLMConfigUpdate):
    """Update LLM configuration. API key stored in memory only."""
    state.llm_config.api_key = config.api_key
    state.llm_config.base_url = config.base_url
    state.llm_config.model = config.model
    state.llm_config.temperature = config.temperature
    state.llm_config.rag_enabled = config.rag_enabled
    return {"success": True, "message": "LLM configuration updated"}


@router.get("/config/llm")
async def get_llm_config():
    """Get current LLM configuration (API key masked)."""
    return {
        "api_key_set": bool(state.llm_config.api_key),
        "api_key_masked": f"{state.llm_config.api_key[:8]}..." if state.llm_config.api_key else "",
        "base_url": state.llm_config.base_url,
        "model": state.llm_config.model,
        "available_models": AVAILABLE_MODELS,
        "temperature": state.llm_config.temperature,
        "rag_enabled": state.llm_config.rag_enabled,
    }


@router.post("/config/llm/verify")
async def verify_llm_config():
    """Verify LLM API key by sending a simple test request."""
    if not state.llm_config.api_key:
        return {"success": False, "error": "API key not configured. Please enter your API key first."}

    try:
        llm = LLMService(
            api_key=state.llm_config.api_key,
            base_url=state.llm_config.base_url,
            model=state.llm_config.model,
            temperature=0.0,
        )
        # Send a minimal test message
        response = llm.chat([
            {"role": "system", "content": "You are a test assistant. Reply with exactly: CONNECTION_OK"},
            {"role": "user", "content": "Test connection. Reply with CONNECTION_OK."},
        ])
        return {
            "success": True,
            "message": "LLM API connection verified successfully!",
            "model": state.llm_config.model,
            "base_url": state.llm_config.base_url,
            "response_preview": response[:200],
        }
    except Exception as e:
        error_msg = str(e)
        return {"success": False, "error": error_msg, "tip": "Check your API key, base URL, and model name."}


# ============================================================
# Custom checklist endpoints
# ============================================================

@router.post("/checklist/upload")
async def upload_checklist(file: UploadFile = File(...)):
    """Upload a custom checklist file (.xlsx or .csv)."""
    content = await file.read()

    items = []
    if file.filename.endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        headers = [cell.value for cell in ws[1]] if ws.max_row > 0 else []

        for row in ws.iter_rows(min_row=2, values_only=True):
            item = {}
            for idx, val in enumerate(row):
                if idx < len(headers) and headers[idx]:
                    item[str(headers[idx])] = str(val) if val else ""
                else:
                    item[f"col_{idx}"] = str(val) if val else ""
            if any(v for v in item.values()):
                items.append(item)
    elif file.filename.endswith(".csv"):
        import csv
        text = content.decode("utf-8")
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            item = {str(k).strip(): str(v).strip() if v else "" for k, v in row.items() if k}
            if any(v for v in item.values()):
                items.append(item)
    else:

        raise HTTPException(status_code=400, detail="Unsupported file format. Use .xlsx or .csv")

    checklist_id = str(uuid.uuid4())[:8]
    checklist = {
        "id": checklist_id,
        "filename": file.filename,
        "items": items,
        "item_count": len(items),
    }
    state.custom_checklists.append(checklist)
    return {"success": True, "checklist_id": checklist_id, "item_count": len(items), "items": items[:10]}


@router.post("/checklist/{checklist_id}/analyze")
async def analyze_checklist(checklist_id: str):
    """Analyze a custom checklist using LLM and provide optimization suggestions."""
    checklist = next((c for c in state.custom_checklists if c["id"] == checklist_id), None)
    if not checklist:
        raise HTTPException(status_code=404, detail="Checklist not found")

    if not state.llm_config.api_key:
        raise HTTPException(status_code=400, detail="LLM API key not configured")

    llm = LLMService(
        api_key=state.llm_config.api_key,
        base_url=state.llm_config.base_url,
        model=state.llm_config.model,
        temperature=state.llm_config.temperature,
    )
    analysis = llm.analyze_custom_checklist(checklist["items"])
    checklist["analysis"] = analysis
    return {"success": True, "analysis": analysis}


# ============================================================
# Assessment endpoints
# ============================================================

class AssessmentRequest(BaseModel):
    task_name: str
    server_ids: List[str]
    benchmarks: List[str]  # CIS, STIG, NIST, PCI — one or more


import threading


def _run_task_background(task: dict, selected_servers: list, req: AssessmentRequest):
    try:
        def log_cb(msg: str):
            task["logs"].append(msg)

        def progress_cb(pct: int):
            task["progress"] = min(99, max(task.get("progress", 0), pct))

        log_cb(f"🚀 Initializing assessment task '{req.task_name}'...")
        progress_cb(10)

        private_key = state.keypair[0]
        ssh_service = SSHService(private_key)
        llm_service = LLMService(
            api_key=state.llm_config.api_key,
            base_url=state.llm_config.base_url,
            model=state.llm_config.model,
            temperature=state.llm_config.temperature,
        )

        if state.llm_config.rag_enabled and not rag_service.is_loaded():
            log_cb("📚 Loading RAG benchmark knowledge base into memory...")
            rag_service.load_benchmarks()

        engine = AssessmentEngine(ssh_service, rag_service, llm_service)

        # Wire up cancellation event
        cancel_event = threading.Event()
        task["_cancel_event"] = cancel_event
        engine.set_cancel_event(cancel_event)

        all_server_results = []
        all_summaries = {}
        total_compliant = 0
        total_non_compliant = 0
        total_unknown = 0
        total_controls = 0
        total_successful = 0
        total_failed = 0

        num_benchmarks = len(req.benchmarks)
        for b_idx, benchmark in enumerate(req.benchmarks, 1):
            if cancel_event.is_set():
                break

            log_cb(f"━━━ Benchmark {b_idx}/{num_benchmarks}: {benchmark} ━━━")

            results = engine.assess_task(
                servers=selected_servers,
                benchmark=benchmark,
                max_workers=min(10, len(selected_servers)),
                log_cb=log_cb,
                progress_cb=progress_cb,
            )

            if results.get("cancelled"):
                task["status"] = "cancelled"
                task["progress"] = task.get("progress", 0)
                task["results"] = results
                log_cb("⏹️ Assessment was cancelled. Partial results saved.")
                return

            # Tag each server result with its benchmark
            for sr in results["server_results"]:
                sr["benchmark"] = benchmark
            all_server_results.extend(results["server_results"])

            s = results["summary"]
            total_compliant += s.get("compliant", 0)
            total_non_compliant += s.get("non_compliant", 0)
            total_unknown += s.get("unknown", 0)
            total_controls += s.get("total_controls", 0)
            total_successful += s.get("successful", 0)
            total_failed += s.get("failed", 0)

        if cancel_event.is_set():
            task["status"] = "cancelled"
            task["progress"] = task.get("progress", 0)
            task["results"] = {"server_results": all_server_results, "summary": all_summaries}
            log_cb("⏹️ Assessment was cancelled. Partial results saved.")
            return

        determined = total_compliant + total_non_compliant
        compliance_rate = f"{(total_compliant / determined * 100 if determined > 0 else 0):.1f}%"

        all_summaries = {
            "total_servers": len(selected_servers),
            "successful": total_successful,
            "failed": total_failed,
            "total_controls": total_controls,
            "compliant": total_compliant,
            "non_compliant": total_non_compliant,
            "unknown": total_unknown,
            "compliance_rate": compliance_rate,
            "benchmarks": req.benchmarks,
        }

        combined_results = {
            "benchmark": ", ".join(req.benchmarks),
            "total_servers": len(selected_servers),
            "server_results": all_server_results,
            "summary": all_summaries,
        }

        log_cb("📝 Generating LLM executive summary & formatted reports...")
        exec_summary = llm_service.generate_report_summary(all_server_results)
        report = report_service.generate_report(
            task_name=req.task_name,
            benchmark=", ".join(req.benchmarks),
            server_results=all_server_results,
            summary=all_summaries,
            executive_summary=exec_summary,
        )

        task["status"] = "completed"
        task["progress"] = 100
        task["results"] = combined_results
        task["report"] = report
        log_cb(f"✅ Task finished! Compliance Rate: {compliance_rate}")
    except Exception as e:
        logger.error(f"Assessment failed: {e}")
        task["status"] = "failed"
        task["error"] = str(e)
        task["logs"].append(f"💥 Task execution error: {str(e)}")


@router.post("/assessment/run")
async def run_assessment(req: AssessmentRequest):
    """Run a security assessment on selected servers in the background with real-time logs."""
    if not state.keypair:
        raise HTTPException(status_code=400, detail="Keypair not generated. Please generate a keypair first.")
    if not state.llm_config.api_key:
        raise HTTPException(status_code=400, detail="LLM API key not configured. Please configure LLM first.")
    if not req.server_ids:
        raise HTTPException(status_code=400, detail="No servers selected for assessment.")
    if not req.benchmarks:
        raise HTTPException(status_code=400, detail="No benchmarks selected for assessment.")

    selected_servers = [s for s in state.servers if s["id"] in req.server_ids]
    if not selected_servers:
        raise HTTPException(status_code=404, detail="Selected servers not found")

    if len(selected_servers) > 10:
        raise HTTPException(status_code=400, detail="Maximum 10 servers per task")

    task_id = str(uuid.uuid4())[:8]
    task = {
        "id": task_id,
        "name": req.task_name,
        "benchmark": ", ".join(req.benchmarks),
        "server_count": len(selected_servers),
        "status": "running",
        "progress": 5,
        "logs": [],
        "results": None,
        "report": None,
    }
    state.tasks.append(task)

    # Start background task execution
    thread = threading.Thread(target=_run_task_background, args=(task, selected_servers, req))
    thread.daemon = True
    thread.start()

    return {
        "success": True,
        "task_id": task_id,
        "message": "Assessment task launched in background",
    }


@router.post("/assessment/cancel/{task_id}")
async def cancel_assessment(task_id: str):
    """Cancel a running assessment task."""
    task = next((t for t in state.tasks if t["id"] == task_id), None)
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")
    if task["status"] not in ("running",):
        raise HTTPException(status_code=400, detail=f"Task is not running (status: {task['status']})")

    cancel_event = task.get("_cancel_event")
    if cancel_event:
        cancel_event.set()

    task["status"] = "cancelling"
    task["logs"].append("⏹️ Cancel requested by user...")
    return {"success": True, "message": "Cancellation requested"}


@router.get("/assessment/status/{task_id}")
async def assessment_status(task_id: str):
    """Get assessment task status and real-time execution logs."""
    task = next((t for t in state.tasks if t["id"] == task_id), None)
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")
    return {
        "id": task["id"],
        "name": task["name"],
        "status": task["status"],
        "progress": task.get("progress", 0),
        "logs": task.get("logs", []),
        "benchmark": task["benchmark"],
        "server_count": task["server_count"],
        "has_report": task.get("report") is not None,
        "summary": task.get("results", {}).get("summary", {}) if task.get("results") else {},
    }



@router.get("/tasks/list")
async def list_tasks():
    """List all assessment tasks."""
    return {
        "tasks": [
            {
                "id": t["id"],
                "name": t["name"],
                "status": t["status"],
                "benchmark": t["benchmark"],
                "server_count": t["server_count"],
                "has_report": t.get("report") is not None,
            }
            for t in state.tasks
        ]
    }


# ============================================================
# Report endpoints
# ============================================================

@router.get("/report/{task_id}/html")
async def get_report_html(task_id: str):
    """Get the HTML report for a task."""
    task = next((t for t in state.tasks if t["id"] == task_id), None)
    if not task or not task.get("report"):
        raise HTTPException(status_code=404, detail="Report not found")
    return HTMLResponse(content=task["report"]["html_content"])


@router.get("/report/{task_id}/markdown")
async def get_report_markdown(task_id: str):
    """Get the Markdown report for a task."""
    task = next((t for t in state.tasks if t["id"] == task_id), None)
    if not task or not task.get("report"):
        raise HTTPException(status_code=404, detail="Report not found")
    return JSONResponse({"markdown": task["report"]["markdown_content"]})


@router.get("/report/{task_id}/download/{format}")
async def download_report(task_id: str, format: str):
    """Download report file (html or md)."""
    task = next((t for t in state.tasks if t["id"] == task_id), None)
    if not task or not task.get("report"):
        raise HTTPException(status_code=404, detail="Report not found")

    if format == "html":
        path = task["report"]["html_path"]
        return FileResponse(path, media_type="text/html", filename=f"{task['name']}_report.html")
    elif format == "md":
        path = task["report"]["markdown_path"]
        return FileResponse(path, media_type="text/markdown", filename=f"{task['name']}_report.md")
    else:
        raise HTTPException(status_code=400, detail="Invalid format. Use 'html' or 'md'.")

